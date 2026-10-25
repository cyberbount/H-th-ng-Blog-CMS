import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report_document():
    doc = docx.Document()

    # Standard A4 Margins: 2.5cm Top/Bottom/Right, 3.0cm Left for binding
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.98)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

        # Cover page carries no header/footer (clean title page)
        section.different_first_page_header_footer = True

        # Header setup
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo cáo Học phần Triển khai & Quản trị Hệ thống Phần mềm — Đề tài 7: Blog CMS")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)

        # Footer setup: institutional line + live page number field
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Trường Đại học Công nghệ Thông tin & Truyền thông (ICTU) — 2026 — Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(148, 163, 184)
        fld = parse_xml(
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr=" PAGE ">'
            '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            '<w:sz w:val="17"/><w:color w:val="94A3B8"/></w:rPr><w:t>2</w:t></w:r></w:fldSimple>'
        )
        p_ftr._p.append(fld)

    return doc

def set_font(run, name="Times New Roman", size=12, color=None, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_font(r, "Times New Roman", 12, RGBColor(30, 41, 59))
    return p

def add_h1(doc, text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, "Times New Roman", 13.5, RGBColor(15, 23, 42), bold=True)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, "Times New Roman", 12, RGBColor(30, 41, 59), bold=True, italic=True)
    return p

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="THÔNG ĐIỆP HỆ THỐNG"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r_title = p.add_run(f"📌 {title}: ")
    set_font(r_title, "Times New Roman", 11.5, RGBColor(37, 99, 235), bold=True)
    r_text = p.add_run(text)
    set_font(r_text, "Times New Roman", 11.5, RGBColor(15, 23, 42))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_text.strip())
    set_font(r, "Consolas", 9.5, RGBColor(241, 245, 249))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(doc, img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption)
        set_font(r_cap, "Times New Roman", 10.5, RGBColor(71, 85, 105), bold=True, italic=True)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[{caption} - Minh chứng hệ thống]")
        set_font(r, "Times New Roman", 10.5, RGBColor(100, 116, 139), italic=True)

def add_styled_table(doc, headers, data, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(caption)
        set_font(r_cap, "Times New Roman", 11, RGBColor(30, 58, 138), bold=True)
        
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Format Header
    hdr_cells = tbl.rows[0].cells
    for idx, title in enumerate(headers):
        hdr_cells[idx].text = title
        set_cell_background(hdr_cells[idx], "1E3A8A")
        set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_font(r, "Times New Roman", 10.5, RGBColor(255, 255, 255), bold=True)
            
    # Format Rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, "Times New Roman", 10, RGBColor(30, 41, 59))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
