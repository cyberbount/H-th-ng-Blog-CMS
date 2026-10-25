import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_core import (
    create_report_document, add_p, add_h1, add_h2, add_h3,
    add_callout, add_code_block, add_figure, add_styled_table, set_font
)

def build():
    print("Generating official Front Matter...")
    doc = create_report_document()
    
    # ==================== TRANG BÌA ====================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(0)
    p_univ.paragraph_format.space_after = Pt(2)
    r = p_univ.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — ĐẠI HỌC THÁI NGUYÊN")
    set_font(r, "Times New Roman", 12, RGBColor(30, 41, 59), bold=True)
    
    p_school = doc.add_paragraph()
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_school.paragraph_format.space_before = Pt(0)
    p_school.paragraph_format.space_after = Pt(2)
    r = p_school.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG")
    set_font(r, "Times New Roman", 13.5, RGBColor(30, 58, 138), bold=True)
    
    p_faculty = doc.add_paragraph()
    p_faculty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_faculty.paragraph_format.space_before = Pt(0)
    p_faculty.paragraph_format.space_after = Pt(20)
    r = p_faculty.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    set_font(r, "Times New Roman", 12.5, RGBColor(30, 58, 138), bold=True)
    
    # Biểu trưng hoặc dấu gạch ngăn
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(30)
    r = p_div.add_run("—————————— 🕮 ——————————")
    set_font(r, "Times New Roman", 12, RGBColor(148, 163, 184))
    
    p_report_title = doc.add_paragraph()
    p_report_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_report_title.paragraph_format.space_before = Pt(10)
    p_report_title.paragraph_format.space_after = Pt(6)
    r = p_report_title.add_run("BÁO CÁO THỰC HÀNH CHUYÊN ĐỀ")
    set_font(r, "Times New Roman", 16, RGBColor(15, 23, 42), bold=True)
    
    p_subject = doc.add_paragraph()
    p_subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subject.paragraph_format.space_before = Pt(0)
    p_subject.paragraph_format.space_after = Pt(24)
    r = p_subject.add_run("MÔN: TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    set_font(r, "Times New Roman", 13.5, RGBColor(71, 85, 105), bold=True)
    
    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.space_before = Pt(10)
    p_topic.paragraph_format.space_after = Pt(40)
    p_topic.paragraph_format.line_spacing = 1.3
    r = p_topic.add_run("ĐỀ TÀI SỐ 7:\nXÂY DỰNG, TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG\nBLOG CMS (LYXLAB RESEARCH PLATFORM)")
    set_font(r, "Times New Roman", 17, RGBColor(29, 78, 216), bold=True)
    
    # Khung thông tin tác giả và giảng viên
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(1.2)
    p_info.paragraph_format.space_before = Pt(10)
    p_info.paragraph_format.space_after = Pt(50)
    p_info.paragraph_format.line_spacing = 1.4
    r = p_info.add_run(
        "Sinh viên thực hiện:\tcyberbount (ICTU)\n"
        "Lớp học phần:\t\tK23A — Quản trị Hệ thống / ATTT\n"
        "Giảng viên hướng dẫn:\tThS. Nguyễn Anh Chuyên\n"
        "Hệ thống triển khai:\t13 Docker Containers (Full Stack)\n"
        "Kho lưu trữ GitHub:\thttps://github.com/cyberbount/H-th-ng-Blog-CMS.git"
    )
    set_font(r, "Times New Roman", 11.5, RGBColor(30, 41, 59))
    
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.paragraph_format.space_before = Pt(20)
    p_loc.paragraph_format.space_after = Pt(0)
    r = p_loc.add_run("Thái Nguyên, Năm 2026")
    set_font(r, "Times New Roman", 11.5, RGBColor(71, 85, 105), italic=True)
    
    # ==================== LỜI CẢM ƠN ====================
    doc.add_page_break()
    p_ack_hdr = doc.add_paragraph()
    p_ack_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ack_hdr.paragraph_format.space_before = Pt(16)
    p_ack_hdr.paragraph_format.space_after = Pt(12)
    r = p_ack_hdr.add_run("LỜI CẢM ƠN")
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    
    add_p(doc, "Em xin bày tỏ lòng biết ơn chân thành và sâu sắc nhất tới thầy Nguyễn Anh Chuyên – Giảng viên hướng dẫn môn học Triển khai và Quản trị Hệ thống Phần mềm (Khoa Công nghệ Thông tin, Trường Đại học Công nghệ Thông tin và Truyền thông – Đại học Thái Nguyên). Trong suốt quá trình học tập và nghiên cứu thực nghiệm, thầy đã tận tình truyền đạt những kiến thức thực tế vô cùng quý báu về kiến trúc ảo hóa container, nguyên lý mạng phân tầng, định tuyến reverse proxy, giám sát đo lường hạ tầng (monitoring) và quản trị nhật ký tập trung (centralized logging).")
    
    add_p(doc, "Học phần Triển khai và Quản trị Hệ thống Phần mềm không chỉ giúp sinh viên tiếp cận góc nhìn của một kỹ sư phát triển phần mềm (Developer) mà còn mở rộng tư duy bao quát của một kỹ sư vận hành hạ tầng (DevOps/SRE) và kỹ sư an toàn hệ thống (SecOps). Nhờ có đề cương học phần chi tiết, các tiêu chí đánh giá minh bạch và những giờ thực hành trực tiếp trên hệ thống, em đã tự tin thiết kế, đóng gói và vận hành trọn vẹn cụm 13 dịch vụ container cho đề tài Blog CMS.")
    
    add_p(doc, "Mặc dù đã nỗ lực tìm hiểu các tài liệu kỹ thuật chuẩn quốc tế của Docker, Prometheus, Grafana, Nginx và Linux Kernel, báo cáo chắc chắn không tránh khỏi những thiếu sót nhất định do giới hạn về mặt thời gian và kiến thức thực tế. Em rất mong nhận được những lời nhận xét, góp ý quý báu từ thầy và các thầy cô trong hội đồng để hệ thống được hoàn thiện hơn nữa trong các giai đoạn phát triển tiếp theo.")
    
    p_ack_sign = doc.add_paragraph()
    p_ack_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ack_sign.paragraph_format.space_before = Pt(12)
    p_ack_sign.paragraph_format.space_after = Pt(12)
    r = p_ack_sign.add_run("Sinh viên thực hiện trân trọng cảm ơn!")
    set_font(r, "Times New Roman", 12, RGBColor(71, 85, 105), italic=True)
    
    # ==================== MỤC LỤC TỔNG QUAN ====================
    doc.add_page_break()
    p_toc_hdr = doc.add_paragraph()
    p_toc_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_toc_hdr.paragraph_format.space_before = Pt(16)
    p_toc_hdr.paragraph_format.space_after = Pt(12)
    r = p_toc_hdr.add_run("MỤC LỤC TỔNG QUAN")
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    
    toc_lines = [
        ('CHƯƠNG 1. GIỚI THIỆU MÔN HỌC VÀ XÁC LẬP DỰ ÁN', '8', True),
        ('1.1 Giới thiệu môn học Triển khai và Quản trị Hệ thống Phần mềm', '8', False),
        ('1.2 Tổng quan đề tài Blog CMS (Nền tảng LyxLab Research)', '9', False),
        ('1.3 Khảo sát nghiệp vụ và nhu cầu sử dụng thực tế', '9', False),
        ('1.4 Xác định vấn đề kỹ thuật và lựa chọn phương án thực hiện', '10', False),
        ('1.5 Xác lập dự án và phạm vi thực hiện', '12', False),
        ('1.6 Yêu cầu đề bài và tiêu chí đánh giá nghiệm thu', '13', False),
        ('CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ ÁP DỤNG', '15', True),
        ('2.1 Mô hình Client - Server và kiến trúc phân tầng Ingress - App - Data', '15', False),
        ('2.2 Công nghệ WordPress CMS, PHP 8.2 và Web Server Apache', '16', False),
        ('2.3 Cơ sở dữ liệu quan hệ MySQL 8.0 và công cụ quản trị phpMyAdmin', '16', False),
        ('2.4 Ảo hóa Container với Docker Engine và Docker Compose v2', '17', False),
        ('2.5 Cổng Nginx Reverse Proxy, SSL SAN và Quản lý lưu lượng', '18', False),
        ('2.6 Hệ thống Giám sát chuỗi thời gian Prometheus TSDB và 5 Exporters', '19', False),
        ('2.7 Thu thập và Phân tích Nhật ký tập trung với Loki 3.0 và Promtail', '20', False),
        ('2.8 Các nguyên lý Gia cố Bảo mật Hệ thống (System Hardening)', '21', False),
        ('CHƯƠNG 3. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG', '22', True),
        ('3.1 Yêu cầu chức năng và phi chức năng của hệ thống Blog CMS', '22', False),
        ('3.2 Xác định Actor và Ma trận phân quyền (RBAC Matrix)', '24', False),
        ('3.3 Biểu đồ Use Case tổng quát và Đặc tả các ca sử dụng tiêu biểu', '25', False),
        ('3.4 Thiết kế Cơ sở dữ liệu quan hệ (Database Schema)', '28', False),
        ('3.5 Thiết kế Phân vùng Mạng và Lưu trữ Bền vững', '30', False),
        ('CHƯƠNG 4. XÂY DỰNG, TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG', '33', True),
        ('4.1 Sơ đồ Kiến trúc Tổng thể Hệ thống do Archify thiết kế', '33', False),
        ('4.2 Thiết kế cấu trúc thư mục và Docker Compose Module hóa', '34', False),
        ('4.3 Quản lý mã nguồn theo 3 mốc Git Commit nghiêm ngặt', '36', False),
        ('4.4 Triển khai Nginx Gateway, Tích hợp SSL SAN và Security Headers', '37', False),
        ('4.5 Cấu hình hệ thống Giám sát Metrics (Prometheus + Grafana Dashboards)', '38', False),
        ('4.6 Cấu hình đường ống thu thập Log tập trung (Promtail + Loki + LogQL)', '40', False),
        ('4.7 Thực thi các biện pháp Bảo mật Hardening thực tế', '41', False),
        ('4.8 Quy trình vận hành, bảo trì, kiểm tra sức khỏe và sao lưu dữ liệu', '41', False),
        ('CHƯƠNG 5. KẾT QUẢ THỰC NGHIỆM VÀ ĐÁNH GIÁ ĐỀ TÀI', '43', True),
        ('5.1 Môi trường thực nghiệm và phương pháp kiểm chứng độc lập', '43', False),
        ('5.2 Kết quả giao diện Blog hiện đại LyxLab (Theme Astra + Design System)', '43', False),
        ('5.3 Kết quả kiểm chứng phân quyền người dùng (Admin, Tác giả, Độc giả)', '44', False),
        ('5.4 Kết quả xuất bản bài viết kỹ thuật và thảo luận phân cấp', '46', False),
        ('5.5 Kết quả kiểm chứng Cơ sở dữ liệu nghiệp vụ trên phpMyAdmin', '48', False),
        ('5.6 Kết quả thu thập Metrics trên Prometheus Server (Targets 100% UP)', '48', False),
        ('5.7 Kết quả trực quan hóa trên 7 Dashboard Grafana tiêu chuẩn', '49', False),
        ('5.8 Kết quả truy vấn LogQL trong Grafana Loki', '50', False),
        ('5.9 Kiểm chứng cô lập mạng và các tiêu chí Bảo mật Hardening', '52', False),
        ('5.10 Bảng đối chiếu nghiệm thu toàn diện theo yêu cầu đề bài', '53', False),
        ('KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN', '56', True),
        ('TÀI LIỆU THAM KHẢO CHÍNH THỨC', '58', True),
        ('Phụ lục A: Hướng dẫn cài đặt và Khởi chạy 1-Click Hệ thống', '59', True),
        ('Phụ lục B: Bảng Tra cứu Cổng Dịch vụ và Thông tin Tài khoản', '59', True),
        ('Phụ lục C: Tập truy vấn PromQL và LogQL phục vụ Kiểm tra', '60', True),
    ]
    
    for title, pg, is_main in toc_lines:
        p_item = doc.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_item.paragraph_format.space_before = Pt(3 if is_main else 1)
        p_item.paragraph_format.space_after = Pt(2 if is_main else 1)
        dots = " ." * int((72 - len(title)) / 2) if len(title) < 70 else "  "
        r_title = p_item.add_run(f"{title} {dots} ")
        set_font(r_title, "Times New Roman", 11.5 if is_main else 11, RGBColor(30, 58, 138) if is_main else RGBColor(51, 65, 85), bold=is_main)
        r_pg = p_item.add_run(pg)
        set_font(r_pg, "Times New Roman", 11.5 if is_main else 11, RGBColor(30, 58, 138) if is_main else RGBColor(51, 65, 85), bold=is_main)
        
    # ==================== DANH MỤC HÌNH VẼ ====================
    doc.add_page_break()
    p_fig_hdr = doc.add_paragraph()
    p_fig_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_fig_hdr.paragraph_format.space_before = Pt(16)
    p_fig_hdr.paragraph_format.space_after = Pt(12)
    r = p_fig_hdr.add_run("DANH MỤC HÌNH VẼ")
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    
    figs = [
        "Hình 3.1: Sơ đồ Use Case tổng quát hệ thống Blog CMS LyxLab",
        "Hình 3.2: Biểu đồ tuần tự (Sequence Diagram) luồng đăng nhập và xác thực phiên",
        "Hình 3.3: Biểu đồ tuần tự luồng biên soạn, xuất bản bài viết kỹ thuật",
        "Hình 3.4: Biểu đồ tuần tự luồng gửi và phản hồi bình luận phân cấp",
        "Hình 3.5: Mô hình quan hệ thực thể (ERD) các bảng nghiệp vụ cốt lõi của WordPress",
        "Hình 4.1: Sơ đồ Kiến trúc Tổng thể Hệ thống Blog CMS do Archify thiết kế",
        "Hình 4.2: Mô hình phân vùng mạng 4 Bridge Network độc lập trong Docker Compose",
        "Hình 4.3: Lịch sử quản lý mã nguồn Git tuân thủ nghiêm ngặt 3 mốc commit",
        "Hình 4.4: Giao diện quản lý cổng Nginx Proxy Manager và cấu hình SSL tự ký",
        "Hình 4.5: Kiến trúc thu thập Metrics đa tầng với Prometheus TSDB và 5 Exporters",
        "Hình 4.6: Đường ống thu thập nhật ký tập trung Promtail cào Docker Socket đẩy về Loki",
        "Hình 5.1: Giao diện trang chủ Cổng thông tin Nghiên cứu Công nghệ LyxLab",
        "Hình 5.2: Giao diện Bảng tin Quản trị của Tác giả (tacgia_k23) bị hạn chế quyền nghiêm ngặt",
        "Hình 5.3: Cấu trúc các bảng dữ liệu WordPress CMS hiển thị trên phpMyAdmin",
        "Hình 5.4: Trạng thái 100% HEALTH UP của toàn bộ các mục tiêu thu thập số liệu trên Prometheus",
        "Hình 5.5: Dashboard tổng quan giám sát tài nguyên máy chủ và container trên Grafana",
        "Hình 5.6: Giao diện Grafana Explore truy vấn nhật ký trực tiếp từ kho lưu trữ Loki",
        "Hình 5.7: Kết quả lọc truy vết các bản ghi lỗi HTTP bằng câu truy vấn LogQL",
        "Hình 5.8: Lệnh docker compose ps xác nhận 13/13 containers đều ở trạng thái Up ổn định"
    ]
    for fig in figs:
        p_f = doc.add_paragraph()
        p_f.paragraph_format.space_before = Pt(2)
        p_f.paragraph_format.space_after = Pt(2)
        r = p_f.add_run(fig)
        set_font(r, "Times New Roman", 11, RGBColor(51, 65, 85))
        
    # ==================== DANH MỤC BẢNG BIỂU ====================
    doc.add_page_break()
    p_tbl_hdr = doc.add_paragraph()
    p_tbl_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_tbl_hdr.paragraph_format.space_before = Pt(16)
    p_tbl_hdr.paragraph_format.space_after = Pt(12)
    r = p_tbl_hdr.add_run("DANH MỤC BẢNG BIỂU")
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    
    tables = [
        "Bảng 1.1: Khảo sát nhu cầu sử dụng của các đối tượng trong hệ thống",
        "Bảng 1.2: So sánh đối chiếu các phương án công nghệ xây dựng Blog CMS",
        "Bảng 1.3: Danh mục phạm vi các sản phẩm bàn giao của đề tài",
        "Bảng 1.4: Đối chiếu tiêu chí đánh giá nghiệm thu theo đề cương học phần",
        "Bảng 2.1: Danh mục các thành phần trong cụm 5 Exporters giám sát",
        "Bảng 3.1: Danh mục các yêu cầu chức năng cốt lõi của hệ thống Blog CMS",
        "Bảng 3.2: Danh mục yêu cầu phi chức năng của hệ thống",
        "Bảng 3.3: Ma trận phân quyền kiểm soát truy cập (RBAC Matrix)",
        "Bảng 3.4: Thiết kế cấu trúc bảng wp_posts lưu trữ bài viết",
        "Bảng 3.5: Thiết kế cấu trúc bảng wp_comments và trường phân cấp comment_parent",
        "Bảng 3.6: Thiết kế 4 phân vùng mạng độc lập trong Docker Compose",
        "Bảng 4.1: Tóm tắt 3 mốc Git Commit triển khai theo quy định của học phần",
        "Bảng 4.2: Các tham số cấu hình Giới hạn tài nguyên phần cứng (Resource Limits)",
        "Bảng 5.1: Danh sách tài khoản kiểm chứng 3 vai trò người dùng trong hệ thống",
        "Bảng 5.2: Danh mục 8 bài viết kỹ thuật thực nghiệm trên hệ thống LyxLab",
        "Bảng 5.3: Bảng đối chiếu tự đánh giá kết quả đề tài theo khung chấm điểm ICTU"
    ]
    for tbl_name in tables:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        r = p_t.add_run(tbl_name)
        set_font(r, "Times New Roman", 11, RGBColor(51, 65, 85))
        
    # ==================== DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT ====================
    doc.add_page_break()
    p_abbr_hdr = doc.add_paragraph()
    p_abbr_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_abbr_hdr.paragraph_format.space_before = Pt(16)
    p_abbr_hdr.paragraph_format.space_after = Pt(12)
    r = p_abbr_hdr.add_run("DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT")
    set_font(r, "Times New Roman", 15.5, RGBColor(30, 58, 138), bold=True)
    
    abbr_data = [
        ["CMS", "Content Management System", "Hệ thống quản trị và phân phối nội dung web."],
        ["RBAC", "Role-Based Access Control", "Mô hình kiểm soát truy cập phân quyền dựa trên vai trò."],
        ["TSDB", "Time-Series Database", "Cơ sở dữ liệu tối ưu hóa cho lưu trữ chuỗi thời gian (Prometheus)."],
        ["SSL / TLS", "Secure Sockets Layer / Transport Layer Security", "Giao thức mật mã hóa bảo đảm an toàn truyền thông trên mạng."],
        ["SAN", "Subject Alternative Name", "Thuộc tính chứng chỉ số hỗ trợ bảo mật cho nhiều tên miền/IP."],
        ["NPM", "Nginx Proxy Manager", "Giao diện trực quan hóa cấu hình Nginx Reverse Proxy và SSL."],
        ["PromQL", "Prometheus Query Language", "Ngôn ngữ truy vấn tính toán số liệu metrics chuỗi thời gian."],
        ["LogQL", "Log Query Language", "Ngôn ngữ truy vấn và tổng hợp dòng nhật ký của Grafana Loki."],
        ["I/O", "Input / Output", "Thao tác đọc ghi dữ liệu vào ổ đĩa hoặc truyền thông mạng."],
        ["RDBMS", "Relational Database Management System", "Hệ quản trị cơ sở dữ liệu quan hệ (MySQL)."],
        ["SRE", "Site Reliability Engineering", "Kỹ thuật bảo đảm độ tin cậy và sẵn sàng của hệ thống."],
        ["DevOps", "Development and Operations", "Văn hóa và phương pháp kết hợp phát triển và vận hành phần mềm."]
    ]
    add_styled_table(doc, ["Thuật ngữ", "Tên đầy đủ tiếng Anh", "Ý nghĩa chuyên môn trong hệ thống"], abbr_data, "Bảng từ viết tắt và thuật ngữ công nghệ cốt lõi")
    
    return doc, add_p, add_h1, add_h2, add_h3, add_callout, add_code_block, add_figure, add_styled_table, set_font

